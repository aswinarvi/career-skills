#!/usr/bin/env python3
"""build_output.py — convert the interview-prep markdown master into md / docx / pdf.

Pure-Python renderers: python-docx for .docx, reportlab for .pdf. If the needed library
is missing, falls back to pandoc when installed, else prints the pip install hint.

Usage:
  python3 build_output.py master.md --format docx --out out/Prep.docx --page a4
  python3 build_output.py master.md --format pdf  --out out/Prep.pdf  --page letter
  python3 build_output.py master.md --format md   --out out/Prep.md

Supported markdown subset (see references/document-structure.md):
  # .. #### headings, paragraphs, - bullets (2-space nesting), 1. numbered lists,
  - [ ] / - [x] checkboxes, | pipe | tables |, ``` fenced code, > blockquotes,
  --- rules, **bold**, *italic*, `code`, [text](url).
"""

import argparse
import os
import re
import shutil
import subprocess
import sys

# ---------------------------------------------------------------- parsing

INLINE_RE = re.compile(
    r"(\*\*.+?\*\*|(?<!\*)\*[^*\n]+?\*(?!\*)|`[^`\n]+?`|\[[^\]\n]+?\]\([^)\s]+?\))"
)
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)\s]+)\)")


def parse_blocks(text):
    """Yield block tuples from markdown text."""
    lines = text.splitlines()
    blocks, para, i = [], [], 0

    def flush_para():
        if para:
            blocks.append(("p", " ".join(s.strip() for s in para)))
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_para()
            lang = stripped[3:].strip()
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(("code", lang, code))
        elif not stripped:
            flush_para()
        elif re.match(r"^#{1,6}\s", stripped):
            flush_para()
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append(("h", min(level, 4), stripped[level:].strip()))
        elif stripped in ("---", "***", "___"):
            flush_para()
            blocks.append(("hr",))
        elif stripped.startswith("|") and stripped.endswith("|"):
            flush_para()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            i -= 1
            if rows:
                width = max(len(r) for r in rows)
                blocks.append(("table", [r + [""] * (width - len(r)) for r in rows]))
        elif stripped.startswith(">"):
            flush_para()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            i -= 1
            blocks.append(("quote", " ".join(q for q in quote if q)))
        else:
            m_ul = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
            m_ol = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
            if m_ul or m_ol:
                flush_para()
                indent = len((m_ul or m_ol).group(1)) // 2
                text_part = m_ul.group(2) if m_ul else m_ol.group(3)
                check = None
                m_chk = re.match(r"^\[( |x|X)\]\s+(.*)$", text_part)
                if m_ul and m_chk:
                    check = m_chk.group(1).lower() == "x"
                    text_part = m_chk.group(2)
                blocks.append(("li", indent, bool(m_ol), check, text_part))
            else:
                para.append(line)
        i += 1
    flush_para()
    return blocks


def tokenize_inline(text):
    """Split text into (kind, payload) tokens: t / b / i / c / link."""
    tokens = []
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            tokens.append(("b", part[2:-2]))
        elif part.startswith("*") and part.endswith("*"):
            tokens.append(("i", part[1:-1]))
        elif part.startswith("`") and part.endswith("`"):
            tokens.append(("c", part[1:-1]))
        else:
            m = LINK_RE.fullmatch(part)
            if m:
                tokens.append(("link", (m.group(1), m.group(2))))
            else:
                tokens.append(("t", part))
    return tokens


def strip_inline(text):
    out = []
    for kind, payload in tokenize_inline(text):
        out.append(payload[0] if kind == "link" else payload)
    return "".join(out)


# ---------------------------------------------------------------- docx

def build_docx(blocks, out_path, page):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    if page == "a4":
        section.page_width, section.page_height = Mm(210), Mm(297)
    else:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Inches(0.75))
    for attr in ("top_margin", "bottom_margin"):
        setattr(section, attr, Inches(0.7))

    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)

    def add_runs(paragraph, text):
        for kind, payload in tokenize_inline(text):
            if kind == "link":
                label, url = payload
                run = paragraph.add_run(label)
                run.font.color.rgb = RGBColor(0x1A, 0x52, 0x76)
                run.underline = True
                if url != label:
                    paragraph.add_run(f" ({url})").font.size = Pt(9)
            else:
                run = paragraph.add_run(payload)
                if kind == "b":
                    run.bold = True
                elif kind == "i":
                    run.italic = True
                elif kind == "c":
                    run.font.name = "Courier New"
                    run.font.size = Pt(9.5)

    first_h1_done = False
    for block in blocks:
        kind = block[0]
        if kind == "h":
            level, text = block[1], strip_inline(block[2])
            if level == 1 and not first_h1_done:
                doc.add_heading(text, 0)
                first_h1_done = True
            else:
                doc.add_heading(text, level)
        elif kind == "p":
            add_runs(doc.add_paragraph(), block[1])
        elif kind == "li":
            _, indent, ordered, check, text = block
            suffix = "" if indent == 0 else f" {min(indent, 2) + 1}"
            style = ("List Number" if ordered else "List Bullet") + suffix
            p = doc.add_paragraph(style=style)
            if check is not None:
                p.add_run("\u2611 " if check else "\u2610 ")
            add_runs(p, text)
        elif kind == "table":
            rows = block[1]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell_text in enumerate(row):
                    cell = table.cell(r, c)
                    cell.paragraphs[0].text = ""
                    add_runs(cell.paragraphs[0], cell_text)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9.5)
                        if r == 0:
                            run.bold = True
            doc.add_paragraph()
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            for n, line in enumerate(block[2]):
                if n:
                    p.add_run().add_break()
                run = p.add_run(line if line.strip() else " ")
                run.font.name = "Courier New"
                run.font.size = Pt(9)
        elif kind == "quote":
            add_runs(doc.add_paragraph(style="Intense Quote"), block[1])
        elif kind == "hr":
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"),
                         ("w:space", "1"), ("w:color", "999999")):
                bottom.set(qn(k), v)
            pBdr.append(bottom)
            pPr.append(pBdr)

    # footer page numbers: "Page X of Y"
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.add_run("Page ")
    for instr in ("PAGE", "NUMPAGES"):
        run = footer_p.add_run()
        for tag, attrs, text in (
            ("w:fldChar", {"w:fldCharType": "begin"}, None),
            ("w:instrText", {"xml:space": "preserve"}, f" {instr} "),
            ("w:fldChar", {"w:fldCharType": "end"}, None),
        ):
            el = OxmlElement(tag)
            for k, v in attrs.items():
                el.set(qn(k), v)
            if text:
                el.text = text
            run._r.append(el)
        if instr == "PAGE":
            footer_p.add_run(" of ")
    for run in footer_p.runs:
        run.font.size = Pt(8.5)

    doc.save(out_path)


# ---------------------------------------------------------------- pdf

def _rl_markup(text):
    def esc(s):
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    out = []
    for kind, payload in tokenize_inline(text):
        if kind == "b":
            out.append(f"<b>{esc(payload)}</b>")
        elif kind == "i":
            out.append(f"<i>{esc(payload)}</i>")
        elif kind == "c":
            out.append(f'<font face="Courier" size="9">{esc(payload)}</font>')
        elif kind == "link":
            label, url = payload
            out.append(f'<link href="{esc(url)}" color="#1a5276"><u>{esc(label)}</u></link>')
        else:
            out.append(esc(payload))
    return "".join(out)


def build_pdf(blocks, out_path, page, title):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (HRFlowable, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle, XPreformatted)

    pagesize = A4 if page == "a4" else letter
    base = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=base["BodyText"], fontSize=10,
                          leading=13.8, spaceAfter=5)
    styles = {
        0: ParagraphStyle("T", parent=base["Title"], fontSize=19, leading=23,
                          spaceAfter=10),
        1: ParagraphStyle("H1", parent=base["Heading1"], fontSize=15.5,
                          spaceBefore=16, spaceAfter=6,
                          textColor=colors.HexColor("#1a3a5c")),
        2: ParagraphStyle("H2", parent=base["Heading2"], fontSize=12.5,
                          spaceBefore=12, spaceAfter=4,
                          textColor=colors.HexColor("#1a3a5c")),
        3: ParagraphStyle("H3", parent=base["Heading3"], fontSize=11,
                          spaceBefore=9, spaceAfter=3),
        4: ParagraphStyle("H4", parent=base["Heading4"], fontSize=10.3,
                          spaceBefore=7, spaceAfter=3),
    }
    code_style = ParagraphStyle("Code", fontName="Courier", fontSize=8.5,
                                leading=10.5, backColor=colors.HexColor("#f4f4f4"),
                                borderPadding=5, spaceAfter=7, leftIndent=4)
    quote_style = ParagraphStyle("Quote", parent=body, leftIndent=16,
                                 textColor=colors.HexColor("#444444"),
                                 borderPadding=3)
    cell_style = ParagraphStyle("Cell", parent=body, fontSize=8.8,
                                leading=11, spaceAfter=0)
    cell_hdr = ParagraphStyle("CellH", parent=cell_style, fontName="Helvetica-Bold")

    doc = SimpleDocTemplate(out_path, pagesize=pagesize,
                            leftMargin=1.7 * cm, rightMargin=1.7 * cm,
                            topMargin=1.6 * cm, bottomMargin=1.7 * cm,
                            title=title or "Interview Prep")
    avail = pagesize[0] - doc.leftMargin - doc.rightMargin
    story, counters, first_h1_done = [], {}, False

    for block in blocks:
        kind = block[0]
        if kind == "h":
            level, text = block[1], _rl_markup(block[2])
            if level == 1 and not first_h1_done:
                story.append(Paragraph(text, styles[0]))
                first_h1_done = True
            else:
                story.append(Paragraph(text, styles[level]))
            counters.clear()
        elif kind == "p":
            story.append(Paragraph(_rl_markup(block[1]), body))
        elif kind == "li":
            _, indent, ordered, check, text = block
            if ordered:
                counters[indent] = counters.get(indent, 0) + 1
                for deeper in [k for k in counters if k > indent]:
                    counters.pop(deeper)
                bullet = f"{counters[indent]}."
            else:
                bullet = {0: "\u2022", 1: "\u25e6"}.get(indent, "\u25aa")
            if check is not None:
                bullet = "\u2611" if check else "\u2610"
            li_style = ParagraphStyle(f"li{indent}", parent=body,
                                      leftIndent=14 + indent * 14,
                                      bulletIndent=4 + indent * 14, spaceAfter=3)
            story.append(Paragraph(_rl_markup(text), li_style, bulletText=bullet))
        elif kind == "table":
            rows = block[1]
            data = [[Paragraph(_rl_markup(c), cell_hdr if r == 0 else cell_style)
                     for c in row] for r, row in enumerate(rows)]
            col_w = avail / len(rows[0])
            tbl = Table(data, colWidths=[col_w] * len(rows[0]), repeatRows=1)
            tbl.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b8b8b8")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9eef3")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.extend([tbl, Spacer(1, 7)])
        elif kind == "code":
            wrapped = []
            for line in block[2]:
                while len(line) > 96:
                    wrapped.append(line[:96])
                    line = "    " + line[96:]
                wrapped.append(line)
            esc = "\n".join(wrapped).replace("&", "&amp;").replace("<", "&lt;")
            story.append(XPreformatted(esc if esc.strip() else " ", code_style))
        elif kind == "quote":
            story.append(Paragraph(_rl_markup(block[1]), quote_style))
        elif kind == "hr":
            story.extend([Spacer(1, 4),
                          HRFlowable(width="100%", thickness=0.6,
                                     color=colors.HexColor("#999999")),
                          Spacer(1, 4)])

    def footer(canvas, docobj):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawCentredString(pagesize[0] / 2, 0.9 * cm,
                                 f"Page {docobj.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


# ---------------------------------------------------------------- fallback + main

def pandoc_fallback(src, out_path, fmt):
    if not shutil.which("pandoc"):
        return False
    cmd = ["pandoc", src, "-o", out_path]
    if fmt == "pdf":
        for engine in ("xelatex", "pdflatex", "wkhtmltopdf", "weasyprint"):
            if shutil.which(engine):
                cmd += [f"--pdf-engine={engine}"]
                break
        else:
            return False
    return subprocess.run(cmd, capture_output=True).returncode == 0


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("input", help="markdown master file")
    ap.add_argument("--format", choices=["md", "docx", "pdf"], required=True)
    ap.add_argument("--out", help="output path (default: input name + extension)")
    ap.add_argument("--page", choices=["a4", "letter"], default="a4")
    ap.add_argument("--title", default=None, help="pdf metadata title")
    args = ap.parse_args()

    if not os.path.isfile(args.input):
        sys.exit(f"error: input not found: {args.input}")
    out_path = args.out or os.path.splitext(args.input)[0] + "." + args.format
    out_dir = os.path.dirname(os.path.abspath(out_path))
    os.makedirs(out_dir, exist_ok=True)

    if args.format == "md":
        if os.path.abspath(args.input) != os.path.abspath(out_path):
            shutil.copyfile(args.input, out_path)
    else:
        with open(args.input, encoding="utf-8") as fh:
            blocks = parse_blocks(fh.read())
        try:
            if args.format == "docx":
                build_docx(blocks, out_path, args.page)
            else:
                title = args.title
                if not title:
                    h1 = next((b[2] for b in blocks if b[0] == "h" and b[1] == 1), None)
                    title = strip_inline(h1) if h1 else None
                build_pdf(blocks, out_path, args.page, title)
        except ImportError as exc:
            lib = "python-docx" if args.format == "docx" else "reportlab"
            print(f"{lib} unavailable ({exc}); trying pandoc fallback...")
            if not pandoc_fallback(args.input, out_path, args.format):
                sys.exit(f"error: install it first: pip install {lib}")

    size_kb = os.path.getsize(out_path) / 1024
    print(f"wrote {out_path} ({size_kb:.0f} KB)")


if __name__ == "__main__":
    main()
